# -*- coding: utf-8 -*-
"""
③ 回填综测总表 · 引擎
============================================================================
依据 ② 产出的《挂科与附加分表.xlsx》，把「是否挂科」「附加分」回填到
各年级「综测成绩表」总表。规则沿用 23 / 24 / 25 级已确认的做法。

── 字段映射 ────────────────────────────────────────────────
  综测总表：A=姓名 B=学号 … E(5)=附加分 … K(11)=是否挂科
  是否挂科 ← 「挂科汇总」学号集合：命中填「是」，否则「否」
  附加分   ← 「附加分汇总」的「100分门数」：无则 0
  注意：E 列参与 F 列公式 =D+E，未获资格也必须填 0，不能留空

── 身份确认（学号 + 班级双重确认，异常必报告）──────────────
  1. 主匹配：12 位学号精确匹配（归一化：strip 并去掉 Excel 文本前缀撇号）
  2. 辅助校验：学号可推导班级，与名单「班级」列比对；不一致 → 「身份存疑」警告
     学号结构 = 年级4位 + 专业码4位 + 班号2位 + 序号2位
     班级名   = 专业名 + 年级后2位 + 专业码前2位 + 班号
     例：202520010501 → 计算机252005（专业码 2001=计算机，班号 05）
     专业码不在映射表内（跟班生如 1608/1609/1802）→ 跳过校验，不误报
  3. 学号非 12 位 → 「学号位数异常」警告并逐条列出

── 年级对齐（避免跨年级误报缺人）──────────────────────────
  源表可能一次含多个年级。回填某个年级目录时，先用该目录学生学号的
  年级众数确定「目标年级」，再只取源表中同年级的名单参与匹配与缺人检查；
  源表无年级列时不过滤（全量匹配）。

── 其他固定约定 ────────────────────────────────────────────
  · 写入前一律先备份到 <目录>/_备份_修改前_时间戳/
  · 一律按名单覆盖（不保留已有值）
  · 「名单有但总表无」属总表缺人，降级为警告，不算填充失败
  · 多 sheet 文件：只处理含表头（姓名+学号+是否挂科）的 sheet

CLI:
  python 回填引擎.py <总表目录> <挂科与附加分表.xlsx> [--dry-run] [--no-backup]
"""
import os
import re
import sys
import shutil
import datetime
from collections import Counter

import openpyxl

ID_LEN = 12

# 专业码(学号第5-8位) → 专业名
MAJOR = {
    '2001': '计算机',
    '2004': '软件工程',
    '2005': '物联网',
    '2006': '智能科学',
    '2007': '大数据',
}

COL_FALLBACK = {'add': 5, 'fail': 11}
SKIP_PREFIX = ('~$', '.~', '._')


# ============================================================
# 基础工具
# ============================================================
def norm_id(v):
    """学号归一化：数字/字符串统一，剥离 Excel 文本前缀撇号与空白。"""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip().lstrip("'").strip()
    return s or None


def id_ok(sid):
    return isinstance(sid, str) and re.fullmatch(r'\d{%d}' % ID_LEN,
                                                 sid) is not None


def derive_klass(sid):
    """由 12 位学号推导班级名；无法推导（专业码未收录/位数不对）返回 None。"""
    if not id_ok(sid):
        return None
    major = sid[4:8]
    if major not in MAJOR:
        return None
    return MAJOR[major] + sid[2:4] + major[:2] + sid[8:10]


# ============================================================
# 读取 ② 产出的源表
# ============================================================
def _hdr_index(ws, row=1):
    out = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row, c).value
        if v is None:
            continue
        out[str(v).strip()] = c
    return out


def load_source(path):
    """读《挂科与附加分表.xlsx》→ (fail_map, add_map, info)

    fail_map: {学号: {'g':年级, 'k':班级}}
    add_map : {学号: (门数, 年级, 班级)}   门数已按人聚合好，不再累加行数
    """
    if not path or not os.path.isfile(path):
        raise RuntimeError('源表不存在：%s' % path)
    wb = openpyxl.load_workbook(path, data_only=True)
    if '挂科汇总' not in wb.sheetnames or '附加分汇总' not in wb.sheetnames:
        raise RuntimeError('源表缺少「挂科汇总」或「附加分汇总」Sheet，'
                           '请确认选的是 ② 产出的《挂科与附加分表.xlsx》')

    fail_map, add_map = {}, {}

    ws = wb['挂科汇总']
    h = _hdr_index(ws)
    c_sid, c_klass, c_g = h.get('学号'), h.get('班级'), h.get('年级')
    if not c_sid:
        raise RuntimeError('「挂科汇总」缺少「学号」列')
    for r in range(2, ws.max_row + 1):
        sid = norm_id(ws.cell(r, c_sid).value)
        if not sid:
            continue
        fail_map[sid] = {
            'g': str(ws.cell(r, c_g).value or '').strip() if c_g else '',
            'k': str(ws.cell(r, c_klass).value or '').strip() if c_klass else '',
        }

    ws = wb['附加分汇总']
    h = _hdr_index(ws)
    c_sid, c_klass, c_g, c_num = (h.get('学号'), h.get('班级'),
                                  h.get('年级'), h.get('100分门数'))
    if not c_sid:
        raise RuntimeError('「附加分汇总」缺少「学号」列')
    for r in range(2, ws.max_row + 1):
        sid = norm_id(ws.cell(r, c_sid).value)
        if not sid:
            continue
        try:
            num = int(ws.cell(r, c_num).value) if c_num else 1
        except (TypeError, ValueError):
            num = 1
        old = add_map.get(sid, (0, '', ''))
        add_map[sid] = (max(old[0], num),
                        str(ws.cell(r, c_g).value or '').strip() if c_g else old[1],
                        str(ws.cell(r, c_klass).value or '').strip()
                        if c_klass else old[2])

    info = {'path': path, 'fail': len(fail_map), 'add': len(add_map),
            'add_cnt': sum(v[0] for v in add_map.values()),
            'grades': sorted({v['g'] for v in fail_map.values()} |
                             {v[1] for v in add_map.values()} - {''})}
    return fail_map, add_map, info


# ============================================================
# 定位总表结构
# ============================================================
def find_header_row(ws):
    """在前 6 行里找表头行：需同时含「姓名」「学号」「是否挂科」。"""
    for r in range(1, min(7, ws.max_row + 1)):
        vals = [str(ws.cell(r, c).value or '').strip()
                for c in range(1, min(ws.max_column, 30) + 1)]
        if '姓名' in vals and '学号' in vals and '是否挂科' in vals:
            return r
    return None


def locate_cols(ws, hdr):
    idx = {}
    for c in range(1, min(ws.max_column, 30) + 1):
        v = ws.cell(hdr, c).value
        if v is not None:
            idx[str(v).strip()] = c
    return {
        'name': idx.get('姓名', 1),
        'sid': idx.get('学号', 2),
        'add': idx.get('附加分', COL_FALLBACK['add']),
        'fail': idx.get('是否挂科', COL_FALLBACK['fail']),
    }


def is_target_xlsx(path):
    """是否含「是否挂科」表头（即综测总表）。只读判断。"""
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception:
        return False
    try:
        for sn in wb.sheetnames:
            if find_header_row(wb[sn]) is not None:
                return True
        return False
    finally:
        try:
            wb.close()
        except Exception:
            pass


def collect_targets(root):
    """自适应：目录下直接是总表 → 单年级；否则逐个年级子目录。

    返回 [(标签, [xlsx路径...]), ...]
    """
    if not root or not os.path.isdir(root):
        raise RuntimeError('目录不存在：%s' % root)

    def files_in(d):
        out = []
        for f in sorted(os.listdir(d)):
            if not f.lower().endswith('.xlsx'):
                continue
            if f.startswith(SKIP_PREFIX) or f == '挂科与附加分表.xlsx':
                continue
            p = os.path.join(d, f)
            if os.path.isfile(p) and is_target_xlsx(p):
                out.append(p)
        return out

    direct = files_in(root)
    if direct:
        return [(os.path.basename(root.rstrip('/')) or root, direct)]

    targets = []
    for name in sorted(os.listdir(root)):
        sub = os.path.join(root, name)
        # 跳过隐藏目录与备份/还原点目录（里面是副本，不能当年级目录处理）
        if not os.path.isdir(sub) or name.startswith('.') \
                or '备份' in name or '还原' in name:
            continue
        fs = files_in(sub)
        if fs:
            targets.append((name, fs))
    return targets


def scan_rows(path):
    """读总表 → [(sheet, 行号, 姓名, 学号, 附加分值, 是否挂科值)]"""
    wb = openpyxl.load_workbook(path, data_only=False)
    rows = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        hdr = find_header_row(ws)
        if hdr is None:
            continue
        cols = locate_cols(ws, hdr)
        for r in range(hdr + 1, ws.max_row + 1):
            nm = ws.cell(r, cols['name']).value
            if nm is None or str(nm).strip() == '':
                continue
            rows.append((sn, r, str(nm).strip(),
                         norm_id(ws.cell(r, cols['sid']).value),
                         ws.cell(r, cols['add']).value,
                         ws.cell(r, cols['fail']).value))
    return rows


def scan_all(targets):
    rows_cache = {}
    for _label, files in targets:
        for p in files:
            if p not in rows_cache:
                try:
                    rows_cache[p] = scan_rows(p)
                except Exception as e:
                    rows_cache[p] = []
                    print('  ⚠ 读取失败已跳过：%s（%s）' % (os.path.basename(p), e))
    return rows_cache


# ============================================================
# 年级对齐 + 计划
# ============================================================
def infer_grade(rows_list):
    """由总表学生学号前 4 位的众数推断该目录所属年级。"""
    c = Counter()
    for _sn, _r, _nm, sid, _a, _f in rows_list:
        if id_ok(sid):
            c[sid[:4]] += 1
    return c.most_common(1)[0][0] if c else None


def build_plans(targets, src, rows_cache):
    """为每个年级目录生成一份「已按年级过滤好的」名单与统计计划。"""
    fail_map, add_map, info = src
    plans = []
    for label, files in targets:
        allrows = [r for f in files for r in rows_cache.get(f, [])]
        g = infer_grade(allrows)

        if g and info['grades']:
            f2 = {k: v for k, v in fail_map.items() if v['g'] in (g, '')}
            a2 = {k: v for k, v in add_map.items() if v[1] in (g, '')}
            dropped = (len(fail_map) - len(f2)) + (len(add_map) - len(a2))
            note = ('已按年级 %s 过滤名单（排除其他年级 %d 条）' % (g, dropped)
                    if dropped else '名单年级 %s' % g)
        else:
            f2, a2 = dict(fail_map), dict(add_map)
            note = '源表无年级列，未做年级过滤（全量匹配）' if not info['grades'] \
                else '未能推断目录年级，未做年级过滤'

        plans.append({'label': label, 'files': files, 'grade': g,
                      'note': note, 'fail': f2, 'add': a2})
    return plans


# ============================================================
# 空跑 / 写入 / 核验
# ============================================================
def phase_dryrun(plans, src):
    _fail, _add, info = src
    print('源表：%s' % info['path'])
    print('  挂科 %d 人 ｜ 附加分 %d 人 / %d 门次 ｜ 涉及年级：%s'
          % (info['fail'], info['add'], info['add_cnt'],
             '、'.join(info['grades']) or '（源表无年级列）'))
    total = {'rows': 0, 'yes': 0, 'no': 0, 'add': 0, 'add_cnt': 0}
    per_file, issues = [], []

    for p in plans:
        print('\n  【%s】目标年级 %s（%s）'
              % (p['label'], p['grade'] or '未识别', p['note']))
        seen = set()
        for path in p['files']:
            rows = rows_of(path, p)
            name = os.path.basename(path)
            stat = {'dir': p['label'], 'file': name, 'rows': 0, 'yes': 0,
                    'no': 0, 'add': 0, 'add_cnt': 0}
            for sn, r, nm, sid, _av, _fv in rows:
                stat['rows'] += 1
                if sid:
                    seen.add(sid)
                is_fail = bool(sid) and sid in p['fail']
                num = p['add'].get(sid, (0, '', ''))[0] if sid else 0
                stat['yes' if is_fail else 'no'] += 1
                if num:
                    stat['add'] += 1
                    stat['add_cnt'] += num
                collect_issues(issues, p, name, r, nm, sid, is_fail)
            per_file.append(stat)
            for k in ('rows', 'yes', 'no', 'add', 'add_cnt'):
                total[k] += stat[k]
            print('     %-28s 学生行=%-4d 是=%-4d 否=%-4d 附加分=%d人/%d分'
                  % (name, stat['rows'], stat['yes'], stat['no'],
                     stat['add'], stat['add_cnt']))

        for sid in sorted(set(p['fail']) - seen):
            issues.append({'kind': '名单有·总表无', 'dir': p['label'],
                           'file': '-', 'row': '-', 'name': p['fail'][sid]['k'],
                           'sid': sid,
                           'msg': '挂科名单有此人（班级%s），总表无对应行，无法写入'
                                  % (p['fail'][sid]['k'] or '-')})
        for sid in sorted(set(p['add']) - seen):
            issues.append({'kind': '名单有·总表无', 'dir': p['label'],
                           'file': '-', 'row': '-', 'name': p['add'][sid][2],
                           'sid': sid,
                           'msg': '附加分名单有此人（班级%s），总表无对应行，无法写入'
                                  % (p['add'][sid][2] or '-')})

    print('\n  合计：学生行 %d ｜ 是=%d ｜ 否=%d ｜ 附加分 %d 人 / %d 门次'
          % (total['rows'], total['yes'], total['no'],
             total['add'], total['add_cnt']))
    return {'per_file': per_file, 'issues': issues, 'total': total}


_ROWS_CACHE = {}


def rows_of(path, plan=None):
    return _ROWS_CACHE.get(path, [])


def collect_issues(issues, p, name, r, nm, sid, is_fail):
    """学号位数异常 + 班级身份校验（仅对命中名单的行做班级校验）。"""
    if not id_ok(sid):
        issues.append({'kind': '学号位数异常', 'dir': p['label'], 'file': name,
                       'row': r, 'name': nm, 'sid': sid,
                       'msg': '总表学号非 %d 位，已按名单保守填「否/0」'
                              % ID_LEN})
        return
    dk = derive_klass(sid)
    if not dk:
        return  # 专业码未收录（跟班生/外专业），跳过校验不误报

    for tag, mp in (('挂科', p['fail']), ('附加分', p['add'])):
        if sid not in mp:
            continue
        g, k = ((mp[sid]['g'], mp[sid]['k']) if tag == '挂科'
                else (mp[sid][1], mp[sid][2]))
        if g and sid[:4] != g:
            # 学号年级 ≠ 所在班级年级 → 跨年级跟班生，属正常，不做班级比对
            issues.append({'kind': '跨年级跟班生', 'dir': p['label'],
                           'file': name, 'row': r, 'name': nm, 'sid': sid,
                           'msg': '学号年级 %s 与所在班级年级 %s 不同，'
                                  '属正常跟班生，已按学号匹配'
                                  % (sid[:4], g)})
        elif k and dk != k:
            issues.append({'kind': '身份存疑', 'dir': p['label'], 'file': name,
                           'row': r, 'name': nm, 'sid': sid,
                           'msg': '%s名单班级「%s」与学号推导班级「%s」不一致'
                                  % (tag, k, dk)})
        break


def backup_plans(plans):
    ts = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    made = []
    for p in plans:
        if not p['files']:
            continue
        d = os.path.dirname(p['files'][0])
        dst = os.path.join(d, '_备份_修改前_%s' % ts)
        os.makedirs(dst, exist_ok=True)
        for f in p['files']:
            shutil.copy2(f, os.path.join(dst, os.path.basename(f)))
        made.append(dst)
        print('  已备份 %d 个文件 → %s' % (len(p['files']), dst))
    return made


def phase_apply(plans, src, do_backup=True):
    """写入并核验。返回 (核验是否通过, 备份目录列表)。"""
    backups = []
    if do_backup:
        print('── 备份 ──')
        backups = backup_plans(plans)

    print('── 写入 ──')
    for p in plans:
        for path in p['files']:
            wb = openpyxl.load_workbook(path, data_only=False)
            n = 0
            for sn in wb.sheetnames:
                ws = wb[sn]
                hdr = find_header_row(ws)
                if hdr is None:
                    continue
                cols = locate_cols(ws, hdr)
                for r in range(hdr + 1, ws.max_row + 1):
                    nm = ws.cell(r, cols['name']).value
                    if nm is None or str(nm).strip() == '':
                        continue
                    sid = norm_id(ws.cell(r, cols['sid']).value)
                    ws.cell(r, cols['fail']).value = (
                        '是' if (sid and sid in p['fail']) else '否')
                    ws.cell(r, cols['add']).value = (
                        p['add'].get(sid, (0, '', ''))[0] if sid else 0)
                    n += 1
            wb.save(path)
            print('  ✔ %s（%s）写入 %d 行' % (os.path.basename(path), p['label'], n))

    print('── 核验 ──')
    ok = phase_verify(plans, src)
    return ok, backups


def phase_verify(plans, src):
    bad, checked = [], 0
    for p in plans:
        for path in p['files']:
            name = os.path.basename(path)
            for sn, r, nm, sid, av, fv in scan_rows(path):
                checked += 1
                wf = '是' if (sid and sid in p['fail']) else '否'
                wa = p['add'].get(sid, (0, '', ''))[0] if sid else 0
                if fv != wf:
                    bad.append('%s 行%s %s 是否挂科 期望%s 实际%r' % (name, r, nm, wf, fv))
                if av != wa:
                    bad.append('%s 行%s %s 附加分 期望%s 实际%r' % (name, r, nm, wa, av))
    if bad:
        print('  ✘ 核验未通过，%d 处不一致（前10）：' % len(bad))
        for b in bad[:10]:
            print('     ' + b)
        return False
    print('  ✔ 核验通过：%d 行逐人回查，是否挂科 / 附加分 与名单完全一致'
          % checked)
    return True


def print_issues(issues):
    """打印异常清单（CLI 与 GUI 共用）。"""
    if not issues:
        return
    print('\n── 需确认的异常（%d 条）──' % len(issues))
    for i in issues:
        print('  · [%s] %s/%s 行%s %s(%s)：%s'
              % (i['kind'], i['dir'], i['file'], i['row'],
                 i['name'] or '-', i['sid'], i['msg']))


# ============================================================
# 报告
# ============================================================
def write_report(out_dir, plans, src, res, backups, verified):
    _f, _a, info = src
    abnormal = [i for i in res['issues'] if i['kind'] != '跨年级跟班生']
    follower = [i for i in res['issues'] if i['kind'] == '跨年级跟班生']
    os.makedirs(out_dir, exist_ok=True)
    p = os.path.join(out_dir, '回填核验报告_%s.md'
                     % datetime.datetime.now().strftime('%Y%m%d_%H%M%S'))
    L = ['# 综测总表回填 · 核验报告', '',
         '生成时间：%s' % datetime.datetime.now().strftime('%Y-%m-%d %H:%M'), '',
         '源表：`%s`（挂科 %d 人 · 附加分 %d 人 / %d 门次）'
         % (info['path'], info['fail'], info['add'], info['add_cnt']), '',
         '## 一、各表结果', '',
         '| 目录 | 文件 | 学生行 | 是否挂科(是) | 否 | 附加分人数 | 附加分总分 |',
         '|---|---|---:|---:|---:|---:|---:|']
    for s in res['per_file']:
        L.append('| %s | %s | %d | %d | %d | %d | %d |'
                 % (s['dir'], s['file'], s['rows'], s['yes'], s['no'],
                    s['add'], s['add_cnt']))
    t = res['total']
    L.append('| **合计** | | **%d** | **%d** | **%d** | **%d** | **%d** |'
             % (t['rows'], t['yes'], t['no'], t['add'], t['add_cnt']))
    L += ['', '## 二、年级对齐', '']
    for pl in plans:
        L.append('- **%s**：目标年级 %s — %s' % (pl['label'], pl['grade'] or '未识别', pl['note']))
    L += ['', '## 三、独立核验', '',
          '✔ 通过：逐人回查，是否挂科 / 附加分 与名单完全一致。' if verified
          else '✘ 未通过，请检查日志中的不一致清单。', '',
          '## 四、需确认的异常（%d 条）' % len(abnormal), '']
    if abnormal:
        L += ['| 类型 | 目录 | 文件 | 行 | 姓名 | 学号 | 说明 |',
              '|---|---|---|---:|---|---|---|']
        for i in abnormal:
            L.append('| %s | %s | %s | %s | %s | `%s` | %s |'
                     % (i['kind'], i['dir'], i['file'], i['row'],
                        i['name'] or '-', i['sid'], i['msg']))
    else:
        L.append('无。')
    if follower:
        L += ['', '## 五、跨年级跟班生（%d 条 · 正常，已按学号匹配）'
              % len(follower), '',
              '| 目录 | 文件 | 行 | 姓名 | 学号 | 说明 |',
              '|---|---|---:|---|---|---|']
        for i in follower:
            L.append('| %s | %s | %s | %s | `%s` | %s |'
                     % (i['dir'], i['file'], i['row'], i['name'] or '-',
                        i['sid'], i['msg']))
    L += ['', '## 六、备份', '']
    for b in backups:
        L.append('- `%s`' % b)
    if not backups:
        L.append('- 本次未备份')
    open(p, 'w', encoding='utf-8').write('\n'.join(L))
    return p


# ============================================================
# 对外主流程（GUI 与 CLI 共用）
# ============================================================
def prepare(root, src_path):
    src = load_source(src_path)
    targets = collect_targets(root)
    if not targets:
        raise RuntimeError('该目录下没有找到综测总表'
                           '（需含「姓名 / 学号 / 是否挂科」表头的 xlsx）')
    rows_cache = scan_all(targets)
    global _ROWS_CACHE
    _ROWS_CACHE = rows_cache
    plans = build_plans(targets, src, rows_cache)
    return src, targets, plans


def main(argv):
    args = [a for a in argv[1:] if not a.startswith('--')]
    dry = '--dry-run' in argv
    nobak = '--no-backup' in argv
    if len(args) < 2:
        print(__doc__)
        return 2
    src, _targets, plans = prepare(args[0], args[1])
    res = phase_dryrun(plans, src)
    print_issues(res['issues'])
    if dry:
        print('\n（空跑模式，未修改任何文件）')
        return 0
    ok, _backups = phase_apply(plans, src, do_backup=not nobak)
    return 0 if ok else 1


if __name__ == '__main__':
    raise SystemExit(main(sys.argv))
